"""
OCR & Text service (v3)
-----------------------

Three OCR engines, chosen automatically or explicitly:

  "basic"    - PyMuPDF's built-in Tesseract integration. Needs only
               Tesseract. Fast, no confidence data.

  "enhanced" - pytesseract + Pillow on the same Tesseract install.
               Gives per-word confidence, PSM control, auto-rotation,
               and (in "thorough" quality mode) tries several image
               preprocessing variants per page and keeps whichever
               gives the highest OCR confidence.

  "trocr"    - Microsoft's TrOCR transformer model (needs torch +
               transformers + numpy), genuinely trained on handwriting
               rather than print. Runs on lines, not whole pages, so
               this module segments each page into text-line bands
               first, then batches those lines through the model.

Handwriting detection:
  - "enhanced" engine: heuristic driven by real OCR confidence.
  - "basic" engine: rough text-shape heuristic.
  - "trocr" engine: handwriting-specialized model.

Speed:
  - "fast" quality mode: one OCR pass per page.
  - "thorough" quality mode: up to 4 preprocessing variants per page.
  - TrOCR lines are batched through the model.
"""

from pathlib import Path
import os
import re
import time
import json
import shutil
import zipfile
import concurrent.futures as cf

import fitz


# ============================================================================
# DIRECTORIES
# ============================================================================

BASE = Path(
    os.path.abspath(
        os.path.join(os.path.dirname(__file__), "..", "temp")
    )
)

OUT = BASE / "ocr_text"
OUT.mkdir(parents=True, exist_ok=True)


# ============================================================================
# LIMITS / SETTINGS
# ============================================================================

MAX_AGE_SECONDS = 2 * 60 * 60

MAX_PAGES = 500

# TrOCR is much slower per page
MAX_PAGES_TROCR = 150

MAX_FILES = 10

NATIVE_TEXT_THRESHOLD = 20

LOW_CONF_CUTOFF = 60

PARALLEL_WORKERS = min(4, os.cpu_count() or 2)

TROCR_BATCH_SIZE = 8


# ============================================================================
# LANGUAGE SETTINGS
# ============================================================================

LANGUAGE_RE = re.compile(
    r"^[a-zA-Z]{3}(\+[a-zA-Z]{3}){0,4}$"
)


LANGUAGES = {
    "eng": "English",
    "fra": "French",
    "deu": "German",
    "spa": "Spanish",
    "ita": "Italian",
    "por": "Portuguese",
    "nld": "Dutch",
    "rus": "Russian",
    "chi_sim": "Chinese (Simplified)",
    "chi_tra": "Chinese (Traditional)",
    "jpn": "Japanese",
    "kor": "Korean",
    "ara": "Arabic",
    "hin": "Hindi",
}


# ============================================================================
# TESSERACT PSM OPTIONS
# ============================================================================

PSM_OPTIONS = {
    3: "Automatic page segmentation (default)",
    4: "Single column of text",
    6: "Single uniform block of text",
    7: "Single text line",
    11: "Sparse text (scattered words)",
}


# ============================================================================
# TrOCR MODELS
# ============================================================================

TROCR_MODELS = {
    "base": "microsoft/trocr-base-handwritten",
    "large": "microsoft/trocr-large-handwritten",
}


# ============================================================================
# OPTIONAL DEPENDENCIES
# ============================================================================

try:
    import docx

    HAS_DOCX = True

except ImportError:
    HAS_DOCX = False


try:
    from PIL import (
        Image,
        ImageOps,
        ImageFilter,
        ImageStat,
    )

    HAS_PIL = True

except ImportError:
    HAS_PIL = False


try:
    import pytesseract
    from pytesseract import Output

    HAS_PYTESSERACT = True

except ImportError:
    HAS_PYTESSERACT = False


try:
    import numpy as np

    HAS_NUMPY = True

except ImportError:
    HAS_NUMPY = False


try:
    import cv2

    HAS_CV2 = True

except ImportError:
    HAS_CV2 = False


try:
    import torch
    from transformers import (
        TrOCRProcessor,
        VisionEncoderDecoderModel,
    )

    HAS_TROCR_DEPS = True

except ImportError:
    HAS_TROCR_DEPS = False


# TrOCR model cache
_trocr_cache = {}


# ============================================================================
# HOUSEKEEPING
# ============================================================================

def cleanup_old_files(max_age=MAX_AGE_SECONDS):
    """
    Remove OCR output files older than max_age seconds.
    """

    now = time.time()

    try:
        for p in OUT.iterdir():

            try:
                age_ok = (
                    now - p.stat().st_mtime
                ) > max_age

                if p.is_file() and age_ok:
                    p.unlink()

                elif p.is_dir() and age_ok:
                    shutil.rmtree(
                        p,
                        ignore_errors=True
                    )

            except OSError:
                continue

    except OSError:
        pass


# ============================================================================
# TESSERACT CAPABILITY CHECK
# ============================================================================

def ocr_probe():
    """
    Check whether Tesseract is available.

    IMPORTANT:
    We intentionally do NOT hard-code a Windows path here.

    Local Windows:
        Tesseract should be installed and available through PATH.

    Render/Linux:
        Tesseract will be installed by the server/Docker environment
        and will also be available through PATH.
    """

    on_path = (
        shutil.which("tesseract") is not None
    )

    probe = None

    try:
        probe = fitz.open()

        page = probe.new_page(
            width=50,
            height=50
        )

        page.get_textpage_ocr(
            flags=0,
            language="eng",
            dpi=72,
            full=True
        )

        return {
            "available": True,
            "on_path": on_path,
            "detail": "",
        }

    except Exception as e:

        detail = (
            str(e).strip()
            or e.__class__.__name__
        )

        if (
            on_path
            and "not found" not in detail.lower()
        ):
            detail = (
                f"{detail} "
                "(tesseract is on PATH, so this is "
                "likely a missing DLL dependency or "
                "a TESSDATA_PREFIX issue, not a missing install)"
            )

        return {
            "available": False,
            "on_path": on_path,
            "detail": detail,
        }

    finally:

        if probe is not None:

            try:
                probe.close()

            except Exception:
                pass


# ============================================================================
# ENHANCED OCR CAPABILITY CHECK
# ============================================================================

def enhanced_probe():
    """
    Check pytesseract + Pillow + Tesseract.
    """

    if not (
        HAS_PIL
        and HAS_PYTESSERACT
    ):
        missing = [
            name
            for name, ok in (
                ("Pillow", HAS_PIL),
                ("pytesseract", HAS_PYTESSERACT),
            )
            if not ok
        ]

        return {
            "available": False,
            "detail": (
                "Missing Python package(s): "
                + ", ".join(missing)
                + "."
            ),
        }

    try:

        pytesseract.get_tesseract_version()

        return {
            "available": True,
            "detail": "",
        }

    except Exception as e:

        return {
            "available": False,
            "detail": str(e),
        }


# ============================================================================
# TrOCR CAPABILITY CHECK
# ============================================================================

def trocr_probe():
    """
    Check whether TrOCR Python dependencies are available.

    The actual model is NOT downloaded here.
    """

    if not (
        HAS_TROCR_DEPS
        and HAS_NUMPY
    ):
        missing = [
            name
            for name, ok in (
                (
                    "torch + transformers",
                    HAS_TROCR_DEPS,
                ),
                (
                    "numpy",
                    HAS_NUMPY,
                ),
            )
            if not ok
        ]

        return {
            "available": False,
            "detail": (
                "Missing Python package(s): "
                + ", ".join(missing)
                + "."
            ),
        }

    device = (
        "cuda"
        if (
            HAS_TROCR_DEPS
            and torch.cuda.is_available()
        )
        else "cpu"
    )

    return {
        "available": True,
        "detail": (
            "ready to load on first use "
            f"(device: {device})"
        ),
    }


# ============================================================================
# CAPABILITIES
# ============================================================================

def capabilities():

    basic = ocr_probe()
    enhanced = enhanced_probe()
    trocr = trocr_probe()

    return {
        "ocr_available": basic["available"],
        "ocr_detail": basic["detail"],

        "enhanced_available": enhanced["available"],
        "enhanced_detail": enhanced["detail"],

        "trocr_available": trocr["available"],
        "trocr_detail": trocr["detail"],

        "trocr_models": TROCR_MODELS,

        "docx_available": HAS_DOCX,

        "languages": LANGUAGES,

        "psm_options": PSM_OPTIONS,

        "max_pages": MAX_PAGES,

        "max_pages_trocr": MAX_PAGES_TROCR,

        "max_files": MAX_FILES,
    }


# ============================================================================
# INPUT NORMALISATION
# ============================================================================

IMAGE_SUFFIXES = {
    ".png",
    ".jpg",
    ".jpeg",
    ".tif",
    ".tiff",
    ".bmp",
    ".webp",
}


def open_as_pdf(input_path):

    suffix = Path(
        input_path
    ).suffix.lower()

    # ------------------------------------------------------------------------
    # PDF
    # ------------------------------------------------------------------------

    if suffix == ".pdf":

        doc = fitz.open(
            str(input_path)
        )

        if doc.needs_pass:

            doc.close()

            raise ValueError(
                "Password-protected PDFs are not supported. "
                "Please remove the password first."
            )

        return doc

    # ------------------------------------------------------------------------
    # IMAGE
    # ------------------------------------------------------------------------

    if suffix in IMAGE_SUFFIXES:

        img_doc = fitz.open(
            str(input_path)
        )

        try:
            rect = img_doc[0].rect

        except Exception:

            img_doc.close()

            raise ValueError(
                "Could not read that image file."
            )

        pdf = fitz.open()

        page = pdf.new_page(
            width=rect.width,
            height=rect.height
        )

        page.insert_image(
            rect,
            filename=str(input_path)
        )

        img_doc.close()

        return pdf

    # ------------------------------------------------------------------------
    # UNSUPPORTED
    # ------------------------------------------------------------------------

    raise ValueError(
        "Unsupported file type. Upload a PDF or an image "
        "(PNG/JPG/TIFF/BMP/WEBP)."
    )


# ============================================================================
# RENDERING / PREPROCESSING
# ============================================================================

def render_page_image(page, dpi):

    pix = page.get_pixmap(
        matrix=fitz.Matrix(
            dpi / 72,
            dpi / 72
        ),
        alpha=False,
    )

    return Image.frombytes(
        "RGB",
        (pix.width, pix.height),
        pix.samples,
    )


def preprocess_image(img):

    try:

        return ImageOps.autocontrast(
            ImageOps.grayscale(img),
            cutoff=1,
        )

    except Exception:

        return img


def _variant_contrast(img):

    return ImageOps.autocontrast(
        ImageOps.grayscale(img),
        cutoff=1,
    )


def _variant_sharpen(img):

    gray = ImageOps.grayscale(img)

    return gray.filter(
        ImageFilter.UnsharpMask(
            radius=2,
            percent=150,
            threshold=3,
        )
    )


def _variant_threshold(img):

    gray = ImageOps.grayscale(img)

    if HAS_CV2 and HAS_NUMPY:

        try:

            arr = np.array(gray)

            _, th = cv2.threshold(
                arr,
                0,
                255,
                cv2.THRESH_BINARY
                + cv2.THRESH_OTSU,
            )

            return Image.fromarray(th)

        except Exception:
            pass

    mean = ImageStat.Stat(
        gray
    ).mean[0]

    return gray.point(
        lambda p: (
            255
            if p > mean
            else 0
        )
    )


# ============================================================================
# BASIC OCR ENGINE
# ============================================================================

def basic_ocr_textpage(
    page,
    language,
    dpi,
):

    try:

        return page.get_textpage_ocr(
            flags=3,
            language=language,
            dpi=dpi,
            full=True,
        )

    except RuntimeError as e:

        raise ValueError(
            "OCR could not run. This server needs "
            "the Tesseract OCR engine installed, "
            "with the tessdata for the language(s) "
            f"you picked. ({e})"
        )

    except Exception as e:

        raise ValueError(
            f"OCR failed on this page: {e}"
        )


def words_from_basic_textpage(tp):

    try:

        raw = tp.extractWORDS()

    except Exception:

        return []

    return [
        (
            w[0],
            w[1],
            w[2],
            w[3],
            w[4],
        )
        for w in raw
        if w[4]
    ]


# ============================================================================
# ENHANCED OCR ENGINE
# ============================================================================

def detect_orientation(
    img,
    min_conf=1.0,
):

    try:

        osd = pytesseract.image_to_osd(
            img,
            output_type=Output.DICT,
        )

        rotate = int(
            osd.get(
                "rotate",
                0
            )
        ) % 360

        conf = float(
            osd.get(
                "orientation_conf",
                0,
            )
            or 0
        )

        if (
            rotate in (
                90,
                180,
                270,
            )
            and conf >= min_conf
        ):
            return rotate

    except Exception:
        pass

    return 0


def _tesseract_image_to_data(
    img,
    language,
    psm,
):

    return pytesseract.image_to_data(
        img,
        lang=language,
        config=f"--psm {int(psm)}",
        output_type=Output.DICT,
    )


def _parse_ocr_data(
    data,
    dpi,
):

    scale = 72.0 / dpi

    words_pt = []

    confs = []

    lines = {}

    n = len(
        data.get(
            "text",
            []
        )
    )

    for i in range(n):

        word = (
            data["text"][i]
            or ""
        ).strip()

        if not word:
            continue

        try:

            conf = int(
                float(
                    data["conf"][i]
                )
            )

        except (
            ValueError,
            TypeError,
        ):

            conf = -1

        if conf >= 0:
            confs.append(conf)

        x0 = (
            data["left"][i]
            * scale
        )

        y0 = (
            data["top"][i]
            * scale
        )

        x1 = (
            data["left"][i]
            + data["width"][i]
        ) * scale

        y1 = (
            data["top"][i]
            + data["height"][i]
        ) * scale

        words_pt.append(
            (
                x0,
                y0,
                x1,
                y1,
                word,
            )
        )

        key = (
            data["block_num"][i],
            data["par_num"][i],
            data["line_num"][i],
        )

        lines.setdefault(
            key,
            []
        ).append(word)

    text = "\n".join(
        " ".join(v)
        for _, v in sorted(
            lines.items()
        )
    ).strip()

    avg_conf = (
        round(
            sum(confs)
            / len(confs),
            1,
        )
        if confs
        else 0.0
    )

    low_conf_ratio = (
        round(
            sum(
                1
                for c in confs
                if c < LOW_CONF_CUTOFF
            )
            / len(confs),
            3,
        )
        if confs
        else 0.0
    )

    return {
        "text": text,
        "avg_confidence": avg_conf,
        "low_conf_ratio": low_conf_ratio,
        "words_pt": words_pt,
    }


def _run_ocr_variant(
    img,
    language,
    psm,
    dpi,
):

    data = _tesseract_image_to_data(
        img,
        language,
        psm,
    )

    return _parse_ocr_data(
        data,
        dpi,
    )


def process_enhanced_page_image(
    base_img,
    language,
    psm,
    dpi,
    enhance_image,
    quality_mode,
):
    """
    Pure function of an already-rendered page image.

    Safe to run from worker threads.
    """

    if quality_mode == "thorough":

        candidates = [
            base_img,
            _variant_contrast(base_img),
            _variant_sharpen(base_img),
            _variant_threshold(base_img),
        ]

        best = None

        for cand in candidates:

            try:

                res = _run_ocr_variant(
                    cand,
                    language,
                    psm,
                    dpi,
                )

            except Exception:

                continue

            if (
                best is None
                or res["avg_confidence"]
                > best["avg_confidence"]
            ):
                best = res

        if best is None:

            raise ValueError(
                "OCR failed on this page across "
                "every preprocessing attempt."
            )

        return best

    img = (
        preprocess_image(base_img)
        if enhance_image
        else base_img
    )

    return _run_ocr_variant(
        img,
        language,
        psm,
        dpi,
    )


# ============================================================================
# TrOCR ENGINE
# ============================================================================

def load_trocr(model_size="base"):

    if model_size in _trocr_cache:

        return _trocr_cache[
            model_size
        ]

    model_name = TROCR_MODELS.get(
        model_size,
        TROCR_MODELS["base"],
    )

    processor = (
        TrOCRProcessor.from_pretrained(
            model_name
        )
    )

    model = (
        VisionEncoderDecoderModel.from_pretrained(
            model_name
        )
    )

    device = (
        "cuda"
        if torch.cuda.is_available()
        else "cpu"
    )

    model.to(device)

    model.eval()

    _trocr_cache[
        model_size
    ] = (
        processor,
        model,
        device,
    )

    return _trocr_cache[
        model_size
    ]


def segment_lines(
    pil_img,
    min_line_height=12,
    gap_threshold=6,
    pad=4,
):
    """
    Classical projection-profile line segmentation.

    Works for single-column pages.
    """

    gray = pil_img.convert("L")

    arr = np.array(gray)

    if HAS_CV2:

        _, binary = cv2.threshold(
            arr,
            0,
            255,
            cv2.THRESH_BINARY_INV
            + cv2.THRESH_OTSU,
        )

    else:

        thresh = (
            arr.mean()
            * 0.85
        )

        binary = (
            arr < thresh
        ).astype(
            "uint8"
        ) * 255

    row_sums = binary.sum(
        axis=1
    )

    ink_rows = (
        row_sums
        > (
            binary.shape[1]
            * 255
            * 0.01
        )
    )

    bands = []

    start = None
    gap = 0

    height = len(
        ink_rows
    )

    for y in range(height):

        if ink_rows[y]:

            if start is None:
                start = y

            gap = 0

        elif start is not None:

            gap += 1

            if gap > gap_threshold:

                end = y - gap

                if (
                    end - start
                    >= min_line_height
                ):
                    bands.append(
                        (
                            max(
                                0,
                                start - pad,
                            ),
                            min(
                                height,
                                end + pad,
                            ),
                        )
                    )

                start = None
                gap = 0

    if start is not None:

        end = height - gap

        if (
            end - start
            >= min_line_height
        ):
            bands.append(
                (
                    max(
                        0,
                        start - pad,
                    ),
                    min(
                        height,
                        end + pad,
                    ),
                )
            )

    return bands


def trocr_generate_batch(
    processor,
    model,
    device,
    line_images,
    batch_size=TROCR_BATCH_SIZE,
    max_length=128,
):

    texts = []

    for start in range(
        0,
        len(line_images),
        batch_size,
    ):

        chunk = [
            im.convert("RGB")
            for im in line_images[
                start:start + batch_size
            ]
        ]

        if not chunk:
            continue

        try:

            pixel_values = (
                processor(
                    images=chunk,
                    return_tensors="pt",
                )
                .pixel_values
                .to(device)
            )

            with torch.no_grad():

                generated_ids = (
                    model.generate(
                        pixel_values,
                        max_length=max_length,
                    )
                )

            decoded = (
                processor.batch_decode(
                    generated_ids,
                    skip_special_tokens=True,
                )
            )

            texts.extend(
                t.strip()
                for t in decoded
            )

        except Exception as e:

            texts.extend(
                [""]
                * len(chunk)
            )

            if start == 0:

                raise ValueError(
                    f"TrOCR inference failed: {e}"
                )

    return texts


def trocr_ocr_page_from_image(
    base_img,
    model_size,
    batch_size=TROCR_BATCH_SIZE,
):

    if not (
        HAS_TROCR_DEPS
        and HAS_NUMPY
    ):

        raise ValueError(
            "The handwriting (TrOCR) engine needs "
            "torch, transformers, and numpy installed."
        )

    try:

        (
            processor,
            model,
            device,
        ) = load_trocr(
            model_size
        )

    except Exception as e:

        raise ValueError(
            "Could not load the TrOCR handwriting model. "
            "It downloads from Hugging Face on first use, "
            "so this needs internet access on this server. "
            f"({e})"
        )

    bands = segment_lines(
        base_img
    )

    if not bands:

        return {
            "text": "",
            "lines": [],
        }

    line_images = [
        base_img.crop(
            (
                0,
                y0,
                base_img.width,
                y1,
            )
        )
        for y0, y1 in bands
    ]

    texts = trocr_generate_batch(
        processor,
        model,
        device,
        line_images,
        batch_size=batch_size,
    )

    lines = [
        {
            "text": t,
            "y0": y0,
            "y1": y1,
        }
        for (
            y0,
            y1,
        ), t in zip(
            bands,
            texts,
        )
        if t
    ]

    return {
        "text": "\n".join(
            ln["text"]
            for ln in lines
        ),
        "lines": lines,
    }


# ============================================================================
# INVISIBLE TEXT HELPERS
# ============================================================================

def embed_invisible_lines(
    page,
    lines,
    dpi,
):

    scale = 72.0 / dpi

    width_pt = page.rect.width

    for ln in lines:

        if not ln["text"]:
            continue

        rect = fitz.Rect(
            0,
            ln["y0"] * scale,
            width_pt,
            ln["y1"] * scale,
        )

        try:

            page.insert_textbox(
                rect,
                ln["text"],
                fontsize=max(
                    6,
                    rect.height * 0.7,
                ),
                fontname="helv",
                render_mode=3,
            )

        except Exception:

            continue


def embed_invisible_text(
    page,
    words_pt,
):

    for (
        x0,
        y0,
        x1,
        y1,
        word,
    ) in words_pt:

        height = max(
            1.0,
            y1 - y0,
        )

        try:

            page.insert_text(
                fitz.Point(
                    x0,
                    y1
                    - height * 0.15,
                ),
                word,
                fontsize=height * 0.85,
                fontname="helv",
                render_mode=3,
            )

        except Exception:

            continue


# ============================================================================
# TEXT HELPERS
# ============================================================================

def native_text(page):

    return (
        page.get_text("text")
        or ""
    ).strip()


def word_count(text):

    return len(
        text.split()
    )


# ============================================================================
# HANDWRITING DETECTION
# ============================================================================

def _bucket(score):

    if score >= 55:
        return "likely handwritten"

    if score >= 30:
        return "possibly handwritten"

    return "likely printed text"


def classify_handwriting_measured(
    avg_conf,
    low_conf_ratio,
):

    score = max(
        0.0,
        min(
            100.0,
            (
                (100 - avg_conf)
                * 0.7
            )
            + (
                low_conf_ratio
                * 100
                * 0.3
            ),
        ),
    )

    return {
        "score": round(
            score,
            1,
        ),
        "label": _bucket(
            score
        ),
        "method": (
            "measured "
            "(OCR confidence)"
        ),
    }


def classify_handwriting_estimated(
    text,
):

    tokens = text.split()

    if not tokens:

        return {
            "score": 50.0,
            "label": (
                "inconclusive "
                "(little/no text detected)"
            ),
            "method": "estimated",
        }

    alpha_ratio = (
        sum(
            c.isalpha()
            for c in text
        )
        / max(
            1,
            len(text),
        )
    )

    short_ratio = (
        sum(
            1
            for t in tokens
            if len(t) <= 2
        )
        / len(tokens)
    )

    score = max(
        0.0,
        min(
            100.0,
            (
                (1 - alpha_ratio)
                * 50
            )
            + (
                short_ratio
                * 50
            ),
        ),
    )

    return {
        "score": round(
            score,
            1,
        ),
        "label": _bucket(
            score
        ),
        "method": (
            "estimated "
            "(rough text-shape heuristic - "
            "install pytesseract + Pillow "
            "for confidence-based detection)"
        ),
    }


def summarize_handwriting(
    page_results,
):

    scored = [
        p["handwriting"]["score"]
        for p in page_results
        if (
            p["mode"] == "ocr"
            and p["handwriting"].get(
                "score"
            )
            is not None
        )
    ]

    if any(
        p["mode"] == "ocr"
        and p["handwriting"].get(
            "method"
        ) == "trocr"
        for p in page_results
    ):

        return {
            "label": (
                "processed with a "
                "handwriting-specialized "
                "model (TrOCR)"
            ),
            "score": None,
        }

    if not scored:

        return {
            "label": (
                "n/a - no OCR was performed "
                "(document has native text)"
            ),
            "score": 0.0,
        }

    avg = (
        sum(scored)
        / len(scored)
    )

    return {
        "label": _bucket(avg),
        "score": round(
            avg,
            1,
        ),
    }


# ============================================================================
# PDF METADATA
# ============================================================================

def document_metadata(doc):

    m = doc.metadata or {}

    return {
        "title": (
            m.get("title")
            or ""
        ),

        "author": (
            m.get("author")
            or ""
        ),

        "creator": (
            m.get("creator")
            or ""
        ),

        "creation_date": _format_pdf_date(
            m.get("creationDate")
        ),

        "page_count": len(doc),
    }


def _format_pdf_date(raw):

    if not raw:
        return ""

    match = re.match(
        r"D:(\d{4})(\d{2})(\d{2})",
        raw,
    )

    return (
        f"{match.group(1)}-"
        f"{match.group(2)}-"
        f"{match.group(3)}"
        if match
        else raw
    )


# ============================================================================
# IMAGE EXTRACTION
# ============================================================================

def extract_embedded_images(
    doc,
    page_numbers,
    dest_dir,
):

    dest_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    seen = set()

    count = 0

    for pno in page_numbers:

        for img in doc[
            pno
        ].get_images(
            full=True
        ):

            xref = img[0]

            if xref in seen:
                continue

            seen.add(xref)

            try:

                base = (
                    doc.extract_image(
                        xref
                    )
                )

                (
                    dest_dir
                    / (
                        f"page{pno + 1}"
                        f"_img{count + 1}."
                        f"{base.get('ext', 'png')}"
                    )
                ).write_bytes(
                    base["image"]
                )

                count += 1

            except Exception:

                continue

    return count


# ============================================================================
# PER-FILE PROCESSING
# ============================================================================

def _process_one(
    input_path,
    settings,
    workdir,
    stem,
):

    language = str(
        settings.get(
            "language"
        )
        or "eng"
    )

    if not LANGUAGE_RE.match(
        language
    ):

        raise ValueError(
            "Invalid language code."
        )

    mode = settings.get(
        "mode",
        "auto",
    )

    if mode not in (
        "auto",
        "ocr",
        "native",
    ):

        raise ValueError(
            "Invalid extraction mode."
        )

    basic_ok = (
        ocr_probe()["available"]
    )

    enhanced_ok = (
        enhanced_probe()["available"]
    )

    engine_pref = settings.get(
        "engine",
        "auto",
    )

    if engine_pref not in (
        "auto",
        "basic",
        "enhanced",
        "trocr",
    ):

        raise ValueError(
            "Invalid OCR engine choice."
        )

    needs_ocr_engine = (
        mode in (
            "auto",
            "ocr",
        )
    )

    if engine_pref == "trocr":

        if not trocr_probe()[
            "available"
        ]:

            raise ValueError(
                "The handwriting (TrOCR) engine "
                "isn't available - install "
                "torch, transformers, and numpy."
            )

        engine = "trocr"

    elif (
        needs_ocr_engine
        and not (
            basic_ok
            or enhanced_ok
        )
    ):

        if mode == "ocr":

            raise ValueError(
                "OCR was requested, but no working "
                "OCR engine was found. Install "
                "Tesseract OCR, or switch to "
                "native text extraction."
            )

        mode = "native"

        engine = "basic"

    else:

        engine = (
            engine_pref
            if engine_pref != "auto"
            else (
                "enhanced"
                if enhanced_ok
                else "basic"
            )
        )

    dpi = max(
        72,
        min(
            400,
            int(
                settings.get(
                    "dpi",
                    200,
                )
            ),
        ),
    )

    psm = int(
        settings.get(
            "psm",
            3,
        )
    )

    if psm not in PSM_OPTIONS:
        psm = 3

    enhance_image = bool(
        settings.get(
            "enhance_image",
            False,
        )
    ) and HAS_PIL

    auto_orient = (
        bool(
            settings.get(
                "auto_orient",
                False,
            )
        )
        and engine == "enhanced"
    )

    quality_mode = settings.get(
        "quality_mode",
        "fast",
    )

    if quality_mode not in (
        "fast",
        "thorough",
    ):
        quality_mode = "fast"

    trocr_model = settings.get(
        "trocr_model",
        "base",
    )

    if trocr_model not in TROCR_MODELS:
        trocr_model = "base"

    rotation = int(
        settings.get(
            "rotation",
            0,
        )
    ) % 360

    if rotation not in (
        0,
        90,
        180,
        270,
    ):
        rotation = 0

    page_filter = settings.get(
        "page_filter",
        "all",
    )

    if page_filter not in (
        "all",
        "odd",
        "even",
    ):
        page_filter = "all"

    want_pdf = bool(
        settings.get(
            "want_searchable_pdf",
            True,
        )
    )

    want_images = bool(
        settings.get(
            "extract_images",
            False,
        )
    )

    # ------------------------------------------------------------------------
    # OPEN DOCUMENT
    # ------------------------------------------------------------------------

    doc = open_as_pdf(
        input_path
    )

    total = len(doc)

    if total == 0:

        doc.close()

        raise ValueError(
            "The document has no pages."
        )

    page_cap = (
        MAX_PAGES_TROCR
        if engine == "trocr"
        else MAX_PAGES
    )

    if total > page_cap:

        doc.close()

        raise ValueError(
            "Too many pages for this engine "
            f"(limit is {page_cap})."
        )

    # ------------------------------------------------------------------------
    # PAGE RANGE
    # ------------------------------------------------------------------------

    start = min(
        max(
            1,
            int(
                settings.get(
                    "page_start",
                    1,
                )
            ),
        ),
        total,
    )

    end = settings.get(
        "page_end"
    )

    end = (
        min(
            total,
            int(end),
        )
        if end
        else total
    )

    if start > end:

        doc.close()

        raise ValueError(
            "Invalid page range."
        )

    def in_range(n):

        if not (
            start <= n <= end
        ):
            return False

        if page_filter == "odd":
            return n % 2 == 1

        if page_filter == "even":
            return n % 2 == 0

        return True

    page_results_map = {}

    processed_indices = []

    enhanced_work = []

    trocr_work = []

    # ------------------------------------------------------------------------
    # FIRST PASS
    # ------------------------------------------------------------------------

    for i in range(total):

        page_number = i + 1

        page = doc[i]

        if rotation:

            page.set_rotation(
                (
                    page.rotation
                    + rotation
                ) % 360
            )

        if not in_range(
            page_number
        ):
            continue

        processed_indices.append(
            i
        )

        text_native = native_text(
            page
        )

        need_ocr = (
            mode == "ocr"
            or (
                mode == "auto"
                and len(text_native)
                < NATIVE_TEXT_THRESHOLD
            )
        )

        # --------------------------------------------------------------------
        # NATIVE TEXT
        # --------------------------------------------------------------------

        if not need_ocr:

            page_results_map[i] = {
                "page": page_number,
                "mode": "native",
                "text": text_native,
                "chars": len(
                    text_native
                ),
                "words": word_count(
                    text_native
                ),
                "confidence": None,
                "orientation_applied": 0,
                "handwriting": {
                    "score": 0.0,
                    "label": (
                        "typed text (native)"
                    ),
                    "method": "native",
                },
            }

            continue

        # --------------------------------------------------------------------
        # BASIC OCR
        # --------------------------------------------------------------------

        if engine == "basic":

            tp = basic_ocr_textpage(
                page,
                language,
                dpi,
            )

            text = (
                page.get_text(
                    "text",
                    textpage=tp,
                )
                or ""
            ).strip()

            if want_pdf:

                embed_invisible_text(
                    page,
                    words_from_basic_textpage(
                        tp
                    ),
                )

            page_results_map[i] = {
                "page": page_number,
                "mode": "ocr",
                "text": text,
                "chars": len(text),
                "words": word_count(text),
                "confidence": None,
                "orientation_applied": 0,
                "handwriting": (
                    classify_handwriting_estimated(
                        text
                    )
                ),
            }

        # --------------------------------------------------------------------
        # ENHANCED OCR
        # --------------------------------------------------------------------

        elif engine == "enhanced":

            orientation_applied = 0

            if auto_orient:

                try:

                    probe_img = (
                        render_page_image(
                            page,
                            min(
                                dpi,
                                150,
                            ),
                        )
                    )

                    orientation_applied = (
                        detect_orientation(
                            probe_img
                        )
                    )

                    if orientation_applied:

                        page.set_rotation(
                            (
                                page.rotation
                                + orientation_applied
                            ) % 360
                        )

                except Exception:

                    orientation_applied = 0

            enhanced_work.append(
                (
                    i,
                    render_page_image(
                        page,
                        dpi,
                    ),
                    orientation_applied,
                )
            )

        # --------------------------------------------------------------------
        # TrOCR
        # --------------------------------------------------------------------

        else:

            trocr_work.append(
                (
                    i,
                    render_page_image(
                        page,
                        dpi,
                    ),
                )
            )

    # =========================================================================
    # ENHANCED OCR PASS
    # =========================================================================

    if enhanced_work:

        with cf.ThreadPoolExecutor(
            max_workers=min(
                PARALLEL_WORKERS,
                len(enhanced_work),
            )
        ) as ex:

            futures = {
                ex.submit(
                    process_enhanced_page_image,
                    img,
                    language,
                    psm,
                    dpi,
                    enhance_image,
                    quality_mode,
                ): (
                    idx,
                    orient,
                )
                for (
                    idx,
                    img,
                    orient,
                ) in enhanced_work
            }

            for fut in cf.as_completed(
                futures
            ):

                idx, orient = futures[
                    fut
                ]

                page = doc[idx]

                try:

                    res = fut.result()

                    if want_pdf:

                        embed_invisible_text(
                            page,
                            res["words_pt"],
                        )

                    page_results_map[idx] = {
                        "page": idx + 1,
                        "mode": "ocr",
                        "text": res["text"],
                        "chars": len(
                            res["text"]
                        ),
                        "words": word_count(
                            res["text"]
                        ),
                        "confidence": (
                            res[
                                "avg_confidence"
                            ]
                        ),
                        "orientation_applied": orient,
                        "handwriting": (
                            classify_handwriting_measured(
                                res[
                                    "avg_confidence"
                                ],
                                res[
                                    "low_conf_ratio"
                                ],
                            )
                        ),
                    }

                except Exception as e:

                    page_results_map[idx] = {
                        "page": idx + 1,
                        "mode": "ocr",
                        "text": "",
                        "chars": 0,
                        "words": 0,
                        "confidence": None,
                        "orientation_applied": orient,
                        "handwriting": {
                            "score": None,
                            "label": (
                                f"OCR failed: {e}"
                            ),
                            "method": "error",
                        },
                    }

    # =========================================================================
    # TrOCR PASS
    # =========================================================================

    for idx, base_img in trocr_work:

        page = doc[idx]

        try:

            result = (
                trocr_ocr_page_from_image(
                    base_img,
                    trocr_model,
                )
            )

            if want_pdf:

                embed_invisible_lines(
                    page,
                    result["lines"],
                    dpi,
                )

            text = result["text"]

            page_results_map[idx] = {
                "page": idx + 1,
                "mode": "ocr",
                "text": text,
                "chars": len(text),
                "words": word_count(text),
                "confidence": None,
                "orientation_applied": 0,
                "handwriting": {
                    "score": None,
                    "label": (
                        "processed with a "
                        "handwriting-specialized "
                        "model (TrOCR)"
                    ),
                    "method": "trocr",
                },
            }

        except ValueError as e:

            page_results_map[idx] = {
                "page": idx + 1,
                "mode": "ocr",
                "text": "",
                "chars": 0,
                "words": 0,
                "confidence": None,
                "orientation_applied": 0,
                "handwriting": {
                    "score": None,
                    "label": (
                        f"TrOCR failed: {e}"
                    ),
                    "method": "error",
                },
            }

    # =========================================================================
    # RESULTS
    # =========================================================================

    page_results = [
        page_results_map[i]
        for i in processed_indices
        if i in page_results_map
    ]

    # -------------------------------------------------------------------------
    # METADATA
    # -------------------------------------------------------------------------

    metadata = document_metadata(
        doc
    )

    metadata[
        "engine_used"
    ] = engine

    # -------------------------------------------------------------------------
    # IMAGE EXTRACTION
    # -------------------------------------------------------------------------

    image_count = (
        extract_embedded_images(
            doc,
            processed_indices,
            workdir
            / f"{stem}_images",
        )
        if want_images
        else 0
    )

    # -------------------------------------------------------------------------
    # COMBINED TEXT
    # -------------------------------------------------------------------------

    combined_text = "\n\n".join(
        f"--- Page {p['page']} ---\n"
        f"{p['text']}"
        for p in page_results
    )

    files = {}

    # -------------------------------------------------------------------------
    # TXT
    # -------------------------------------------------------------------------

    txt_path = (
        workdir
        / f"{stem}.txt"
    )

    txt_path.write_text(
        combined_text,
        encoding="utf-8",
    )

    files["txt"] = (
        f"{stem}.txt"
    )

    # -------------------------------------------------------------------------
    # JSON
    # -------------------------------------------------------------------------

    json_path = (
        workdir
        / f"{stem}.json"
    )

    json_path.write_text(
        json.dumps(
            page_results,
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    files["json"] = (
        f"{stem}.json"
    )

    # -------------------------------------------------------------------------
    # MARKDOWN
    # -------------------------------------------------------------------------

    md = "\n\n".join(
        f"## Page {p['page']}\n\n"
        f"{p['text']}"
        for p in page_results
    )

    md_path = (
        workdir
        / f"{stem}.md"
    )

    md_path.write_text(
        md,
        encoding="utf-8",
    )

    files["md"] = (
        f"{stem}.md"
    )

    # -------------------------------------------------------------------------
    # DOCX
    # -------------------------------------------------------------------------

    if HAS_DOCX:

        try:

            d = docx.Document()

            for p in page_results:

                d.add_heading(
                    f"Page {p['page']}",
                    level=2,
                )

                d.add_paragraph(
                    p["text"]
                    or "(no text found)"
                )

            d.save(
                str(
                    workdir
                    / f"{stem}.docx"
                )
            )

            files["docx"] = (
                f"{stem}.docx"
            )

        except Exception:

            pass

    # -------------------------------------------------------------------------
    # SEARCHABLE PDF
    # -------------------------------------------------------------------------

    if want_pdf:

        pdf_name = (
            f"{stem}.pdf"
        )

        doc.save(
            str(
                workdir
                / pdf_name
            ),
            garbage=4,
            deflate=True,
            clean=True,
        )

        files["pdf"] = pdf_name

    # -------------------------------------------------------------------------
    # CLOSE DOCUMENT
    # -------------------------------------------------------------------------

    doc.close()

    # =========================================================================
    # STATISTICS
    # =========================================================================

    stats = {
        "pages_total": total,

        "pages_processed": len(
            page_results
        ),

        "pages_ocrd": sum(
            1
            for p in page_results
            if p["mode"] == "ocr"
        ),

        "pages_native": sum(
            1
            for p in page_results
            if p["mode"] == "native"
        ),

        "total_words": sum(
            p["words"]
            for p in page_results
        ),

        "total_chars": sum(
            p["chars"]
            for p in page_results
        ),

        "images_extracted": image_count,
    }

    # =========================================================================
    # RETURN
    # =========================================================================

    return {
        "pages": page_results,

        "stats": stats,

        "metadata": metadata,

        "handwriting_summary": (
            summarize_handwriting(
                page_results
            )
        ),

        "files": files,
    }


# ============================================================================
# BATCH ENTRY POINT
# ============================================================================

def process(
    input_paths,
    settings,
    source_names=None,
):

    cleanup_old_files()

    if not input_paths:

        raise ValueError(
            "No files to process."
        )

    if len(input_paths) > MAX_FILES:

        raise ValueError(
            f"Too many files at once "
            f"(limit is {MAX_FILES})."
        )

    token = (
        os.urandom(16).hex()
    )

    workdir = (
        OUT / token
    )

    workdir.mkdir(
        parents=True,
        exist_ok=True,
    )

    results = []

    used_stems = set()

    for idx, path in enumerate(
        input_paths
    ):

        name = (
            source_names[idx]
            if (
                source_names
                and idx < len(source_names)
            )
            else path.stem
        )

        stem = re.sub(
            r"[^A-Za-z0-9_-]+",
            "_",
            name,
        ).strip("_") or (
            f"document_{idx + 1}"
        )

        base_stem = stem

        n = 1

        while stem in used_stems:

            n += 1

            stem = (
                f"{base_stem}_{n}"
            )

        used_stems.add(stem)

        try:

            result = _process_one(
                path,
                settings,
                workdir,
                stem,
            )

            result["filename"] = name

            result["error"] = None

        except ValueError as e:

            result = {
                "filename": name,
                "error": str(e),
                "pages": [],
                "stats": {},
                "files": {},
            }

        results.append(
            result
        )

    # =========================================================================
    # ZIP RESULT
    # =========================================================================

    zip_path = (
        OUT
        / f"{token}.zip"
    )

    with zipfile.ZipFile(
        zip_path,
        "w",
        zipfile.ZIP_DEFLATED,
    ) as zf:

        for p in workdir.rglob("*"):

            if p.is_file():

                zf.write(
                    p,
                    arcname=p.relative_to(
                        workdir
                    ),
                )

    return token, results


# ============================================================================
# DOWNLOAD / RESULT HELPERS
# ============================================================================

def workfile_path(
    token,
    filename,
):

    if (
        not token
        or len(token) != 32
        or any(
            c not in
            "0123456789abcdef"
            for c in token.lower()
        )
    ):

        return None

    safe_name = re.sub(
        r"[^A-Za-z0-9_.-]+",
        "_",
        filename or "",
    )

    p = (
        OUT
        / token
        / safe_name
    ).resolve()

    workdir = (
        OUT
        / token
    ).resolve()

    if (
        workdir not in p.parents
        and p != workdir
    ):

        return None

    return (
        p
        if p.is_file()
        else None
    )


def zip_path_for(
    token,
):

    if (
        not token
        or len(token) != 32
        or any(
            c not in
            "0123456789abcdef"
            for c in token.lower()
        )
    ):

        return None

    p = (
        OUT
        / f"{token}.zip"
    ).resolve()

    return (
        p
        if p.is_file()
        else None
    )